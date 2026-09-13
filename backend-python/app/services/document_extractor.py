"""Extract plain text from PRD documents (PDF, DOCX, TXT, MD)."""

from __future__ import annotations

from pathlib import Path
from typing import Optional

from app.core.exceptions import ValidationFailedError
from app.core.logging import get_logger

logger = get_logger(__name__)

SUPPORTED_PRD_EXTENSIONS = {"pdf", "docx", "txt", "md", "text"}


def _extension(filename: str) -> str:
    return Path(filename).suffix.lower().lstrip(".")


def extract_text_from_bytes(filename: str, data: bytes) -> str:
    """Extract UTF-8 text from supported document bytes."""
    ext = _extension(filename)
    if ext not in SUPPORTED_PRD_EXTENSIONS:
        raise ValidationFailedError(
            f"Unsupported PRD file type '.{ext}'",
            details={"filename": filename, "allowed": sorted(SUPPORTED_PRD_EXTENSIONS)},
        )
    if not data:
        raise ValidationFailedError("Uploaded PRD file is empty", details={"filename": filename})

    if ext in {"txt", "md", "text"}:
        return _extract_text(data)
    if ext == "pdf":
        return _extract_pdf(data, filename)
    if ext == "docx":
        return _extract_docx(data, filename)
    raise ValidationFailedError(f"Unsupported PRD file type '.{ext}'")


def extract_text_from_path(path: Path, *, original_filename: Optional[str] = None) -> str:
    name = original_filename or path.name
    data = path.read_bytes()
    return extract_text_from_bytes(name, data)


def _extract_text(data: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = data.decode("utf-8", errors="replace")
    cleaned = text.strip()
    if not cleaned:
        raise ValidationFailedError("PRD text file has no readable content")
    return cleaned


def _extract_pdf(data: bytes, filename: str) -> str:
    try:
        from io import BytesIO

        from pypdf import PdfReader
    except ImportError as exc:
        raise ValidationFailedError(
            "PDF support requires pypdf. Install with: pip install pypdf",
            details={"filename": filename},
        ) from exc

    try:
        reader = PdfReader(BytesIO(data))
        pages: list[str] = []
        for i, page in enumerate(reader.pages):
            try:
                page_text = page.extract_text() or ""
            except Exception as exc:  # noqa: BLE001 - page-level resilience
                logger.warning("Failed to extract PDF page %s from %s: %s", i, filename, exc)
                page_text = ""
            if page_text.strip():
                pages.append(page_text.strip())
        text = "\n\n".join(pages).strip()
    except Exception as exc:  # noqa: BLE001
        raise ValidationFailedError(
            f"Failed to read PDF '{filename}': {exc}",
            details={"filename": filename},
        ) from exc

    if not text:
        raise ValidationFailedError(
            f"No extractable text found in PDF '{filename}' (may be scanned/image-only)",
            details={"filename": filename},
        )
    return text


def _extract_docx(data: bytes, filename: str) -> str:
    try:
        from io import BytesIO

        from docx import Document
    except ImportError as exc:
        raise ValidationFailedError(
            "DOCX support requires python-docx. Install with: pip install python-docx",
            details={"filename": filename},
        ) from exc

    try:
        doc = Document(BytesIO(data))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
                if cells:
                    paragraphs.append(" | ".join(cells))
        text = "\n".join(paragraphs).strip()
    except Exception as exc:  # noqa: BLE001
        raise ValidationFailedError(
            f"Failed to read DOCX '{filename}': {exc}",
            details={"filename": filename},
        ) from exc

    if not text:
        raise ValidationFailedError(
            f"No extractable text found in DOCX '{filename}'",
            details={"filename": filename},
        )
    return text
